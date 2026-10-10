"""Импорт поисков из выгрузки Авито (xlsx): парсинг, staging в Redis, сборка поиска."""

import io
import json
import re
import uuid
from dataclasses import dataclass
from urllib.parse import quote

from openpyxl import load_workbook
from redis.asyncio import Redis

MAX_ROWS = 1000
STAGING_KEY = "file-import:{token}"
STAGING_TTL_SECONDS = 2 * 3600
MAX_QUERY_WORDS = 5

COLOR_WORDS = (
    "black",
    "white",
    "blue",
    "red",
    "green",
    "grey",
    "gray",
    "silver",
    "gold",
    "copper",
    "anthracite",
    "aluminium",
    "aluminum",
    "natural",
    "titanium",
    "midnight",
    "starlight",
    "purple",
    "yellow",
    "orange",
    "pink",
    "brown",
    "beige",
    "цвет",
)
NOISE_WORDS = ("новый", "новое", "новые", "продам", "продаю", "оригинал", "гарантия")


@dataclass(frozen=True, slots=True)
class FileRow:
    avito_id: int | None
    title: str
    price: float | None
    status: str | None
    category: str | None

    def to_dict(self) -> dict[str, object]:
        return {
            "avito_id": self.avito_id,
            "title": self.title,
            "price": self.price,
            "status": self.status,
            "category": self.category,
        }


def _as_float(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, int | float):
        return float(value) if value else None
    digits = re.sub(r"[^\d]", "", str(value))
    return float(digits) if digits else None


def _as_int(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    digits = re.sub(r"[^\d]", "", str(value))
    return int(digits) if digits else None


def parse_xlsx(data: bytes) -> list[FileRow]:
    """Читает первый лист выгрузки Авито: Id/AvitoId, Title, Price, AvitoStatus, Category."""
    # Без read_only: у части выгрузок Авито битые метаданные размеров листа,
    # и read_only-режим видит только первую строку.
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    index: dict[str, int] = {}
    for position, name in enumerate(header):
        if name is None:
            continue
        index[str(name).strip().lower()] = position

    def column(*names: str) -> int | None:
        for name in names:
            if name in index:
                return index[name]
        return None

    id_col = column("avitoid", "id")
    title_col = column("title", "название", "заголовок")
    price_col = column("price", "цена")
    status_col = column("avitostatus", "статус")
    category_col = column("category", "категория")
    if title_col is None:
        return []

    collected: list[FileRow] = []
    for row in rows:
        if len(collected) >= MAX_ROWS:
            break
        title = str(row[title_col] or "").strip() if len(row) > title_col else ""
        avito_id = _as_int(row[id_col]) if id_col is not None and len(row) > id_col else None
        if not title and avito_id is None:
            continue
        collected.append(
            FileRow(
                avito_id=avito_id,
                title=title,
                price=_as_float(row[price_col])
                if price_col is not None and len(row) > price_col
                else None,
                status=(
                    str(row[status_col]).strip()
                    if status_col is not None and len(row) > status_col and row[status_col]
                    else None
                ),
                category=(
                    str(row[category_col]).strip()
                    if category_col is not None and len(row) > category_col and row[category_col]
                    else None
                ),
            )
        )
    return collected


def fallback_query(title: str) -> str:
    """Простой запрос без AI: слова без цветов/шума, максимум MAX_QUERY_WORDS."""
    words = re.findall(r"[0-9A-Za-zА-Яа-яЁё&+.-]+", title)
    cleaned = [
        word
        for word in words
        if word.lower() not in COLOR_WORDS and word.lower() not in NOISE_WORDS
    ]
    return " ".join(cleaned[:MAX_QUERY_WORDS]) or title[:80]


def validate_generated_filter(
    *, title: str, query: str, keyword_groups: list[list[str]], exclude_keywords: list[str]
) -> bool:
    """Reject unsafe LLM filters before they become scheduled searches."""
    title_tokens = set(re.findall(r"[0-9a-zа-яё]+", title.lower()))
    query_tokens = set(re.findall(r"[0-9a-zа-яё]+", query.lower()))
    if not query.strip() or not query_tokens & title_tokens:
        return False
    if not 1 <= len(keyword_groups) <= 3 or any(not group for group in keyword_groups):
        return False
    include_tokens = {
        token.lower() for group in keyword_groups for value in group for token in value.split()
    }
    if any(
        not any(set(re.findall(r"[0-9a-zа-яё]+", value.lower())) & title_tokens for value in group)
        for group in keyword_groups
    ):
        return False
    exclude_tokens = {token.lower() for value in exclude_keywords for token in value.split()}
    if include_tokens & exclude_tokens:
        return False
    return not any(len(value.split()) > 6 for value in exclude_keywords)


def search_params(
    *,
    query: str,
    keyword_groups: list[list[str]] | None,
    exclude_keywords: list[str] | None,
    regions: list[str] | None = None,
    exclude_regions: list[str] | None = None,
) -> dict[str, object]:
    params: dict[str, object] = {
        "keyword_groups": keyword_groups or [[query]] if query else [],
        "exclude_keywords": exclude_keywords or [],
    }
    if regions:
        params["regions"] = regions
    if exclude_regions:
        params["exclude_regions"] = exclude_regions
    return params


def search_url(query: str, city: str | None = None) -> str:
    section = city or "all"
    return f"https://www.avito.ru/{section}?q={quote(query)}"


async def save_staging(redis: Redis, rows: list[FileRow]) -> str:
    token = uuid.uuid4().hex
    payload = json.dumps([row.to_dict() for row in rows], ensure_ascii=False)
    await redis.set(STAGING_KEY.format(token=token), payload, ex=STAGING_TTL_SECONDS)
    return token


async def load_staging(redis: Redis, token: str) -> list[dict[str, object]] | None:
    raw = await redis.get(STAGING_KEY.format(token=token))
    if not raw:
        return None
    try:
        data = json.loads(raw)
    except (TypeError, ValueError):
        return None
    return data if isinstance(data, list) else None
